#!/usr/bin/env python3
"""
export_xlsx_v2.py — FASE 2, arquivo NOVO. Lê .data/social-content-v2-export.json
(gerado por src/v2/export-v2.js) e escreve reports/social-content/social_content_v2.xlsx
com duas abas: PECAS (uma linha por peça) e CALENDARIO (ordenado por data/horário).

Não gera nenhuma imagem/vídeo, não toca em 313_pecas_smart_pet_gadgets.xlsx (V1).
Reaproveita o mesmo estilo visual (cabeçalho colorido, largura de coluna ajustada,
wrap de texto) usado em generate_assets.py, sem importar esse módulo (evita
qualquer efeito colateral do gerador de mídia V1).
"""
import json
import sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl não está instalado. Rode: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[2]
EXPORT_JSON = ROOT / ".data" / "social-content-v2-export.json"
XLSX_OUT = ROOT / "reports" / "social-content" / "social_content_v2.xlsx"

HEADER_FILL = PatternFill(start_color="FF6A2C", end_color="FF6A2C", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

PECAS_COLUMNS = [
    ("ID", "ID"),
    ("DATA_SUGERIDA", "Data sugerida"),
    ("DIA", "Dia"),
    ("HORARIO_SUGERIDO", "Horário"),
    ("CLUSTER", "Cluster"),
    ("ARTIGO", "Artigo"),
    ("ARTICLE_URL", "URL"),
    ("FORMATO", "Formato"),
    ("TIPO", "Tipo"),
    ("OBJETIVO", "Objetivo"),
    ("HOOK", "Hook"),
    ("LEGENDA", "Legenda"),
    ("CTA", "CTA"),
    ("HASHTAGS", "Hashtags"),
    ("ARQUIVO_MIDIA", "Arquivo mídia (planejado)"),
    ("STATUS", "Status"),
    ("SCORE", "Score"),
    ("OBSERVACAO", "Observação"),
]

CALENDARIO_COLUMNS = [
    ("DATA_SUGERIDA", "Data"),
    ("DIA", "Dia"),
    ("HORARIO_SUGERIDO", "Horário"),
    ("ID", "ID"),
    ("FORMATO", "Formato"),
    ("TIPO", "Tipo"),
    ("ARTIGO", "Artigo"),
    ("STATUS", "Status"),
]


def build_sheet(wb, name, columns, rows, sort_key=None):
    ws = wb.create_sheet(name)
    ws.append([label for _, label in columns])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    data = sorted(rows, key=sort_key) if sort_key else rows
    for row in data:
        values = []
        for key, _ in columns:
            v = row.get(key)
            if isinstance(v, list):
                v = " ".join(v)
            values.append(v)
        ws.append(values)
    for i, (key, _) in enumerate(columns, start=1):
        width = 40 if key in ("LEGENDA", "OBSERVACAO", "HOOK") else 18
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    return ws


def main():
    if not EXPORT_JSON.exists():
        print(f"Não encontrado: {EXPORT_JSON} — rode export-v2.js primeiro.", file=sys.stderr)
        sys.exit(1)
    data = json.loads(EXPORT_JSON.read_text(encoding="utf-8"))
    rows = data["rows"]

    wb = Workbook()
    wb.remove(wb.active)
    build_sheet(wb, "PECAS", PECAS_COLUMNS, rows, sort_key=lambda r: r.get("ID", ""))
    build_sheet(
        wb, "CALENDARIO", CALENDARIO_COLUMNS, rows,
        sort_key=lambda r: (r.get("DATA_SUGERIDA") or "", r.get("HORARIO_SUGERIDO") or ""),
    )

    XLSX_OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(XLSX_OUT)
    print(f"OK: {len(rows)} peças -> {XLSX_OUT}")


if __name__ == "__main__":
    main()
